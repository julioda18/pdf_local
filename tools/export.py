"""Exportación de PDF a formatos de Office editables: Word, Excel y PowerPoint."""

import importlib.util
import logging
from io import BytesIO

from core.errors import ToolError
from core.pdf_io import open_pdf

SLIDE_DPI = 150
SLIDE_JPEG_QUALITY = 85
POINTS_TO_EMU = 12700
MAX_SHEET_TITLE = 31
MAX_COLUMN_WIDTH = 60


def _installed(*modules):
    return all(importlib.util.find_spec(m) is not None for m in modules)


def word_available():
    """``True`` si pdf2docx está instalado."""
    return _installed("pdf2docx")


def excel_available():
    """``True`` si openpyxl está instalado."""
    return _installed("openpyxl")


def ppt_available():
    """``True`` si python-pptx está instalado."""
    return _installed("pptx")


def pdf_to_word(inputs, out, **_):
    """Reconstruye texto, tablas e imágenes del PDF como un ``.docx`` editable.

    Un PDF escaneado sale como imágenes: conviene pasarlo antes por OCR.
    """
    # pdf2docx importa la API antigua "fitz" y avisa al importarse; se carga solo al usarla.
    from pdf2docx import Converter

    source = inputs[0]
    # Valida antes (dañado / con contraseña) para dar el mensaje claro de open_pdf.
    open_pdf(source).close()
    result = out.with_suffix(".docx")
    logging.getLogger("pdf2docx").setLevel(logging.WARNING)
    converter = Converter(str(source))
    try:
        converter.convert(str(result))
    except Exception as error:
        raise ToolError(f"No se pudo convertir {source.name} a Word: {error}") from None
    finally:
        converter.close()
    return result


def pdf_to_excel(inputs, out, **_):
    """Vuelca cada tabla detectada en una hoja de un ``.xlsx``.

    Si el PDF no tiene tablas reconocibles, crea una hoja por página con su texto
    línea a línea, para que el resultado nunca salga vacío.
    """
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    with open_pdf(inputs[0]) as doc:
        sheets = list(_table_sheets(doc)) or list(_text_sheets(doc))
    if not sheets:
        raise ToolError("El PDF no tiene texto: pásalo antes por OCR")
    for title, rows in sheets:
        sheet = workbook.create_sheet(title[:MAX_SHEET_TITLE])
        for row in rows:
            sheet.append(row)
        _fit_columns(sheet)
    result = out.with_suffix(".xlsx")
    workbook.save(result)
    return result


def _table_sheets(doc):
    """``(título, filas)`` por cada tabla de cada página."""
    for number, page in enumerate(doc, 1):
        for index, table in enumerate(page.find_tables().tables, 1):
            rows = _drop_empty_columns([[_clean(cell) for cell in row] for row in table.extract()])
            if rows:
                yield f"Pág {number} - Tabla {index}", rows


def _drop_empty_columns(rows):
    """Quita columnas sin contenido, que aparecen cuando el PDF dibuja bordes dobles."""
    keep = [c for c in range(max(map(len, rows), default=0)) if any(c < len(r) and r[c] for r in rows)]
    return [[row[c] if c < len(row) else "" for c in keep] for row in rows] if keep else []


def _text_sheets(doc):
    """``(título, filas)`` por página con texto, una línea por fila."""
    for number, page in enumerate(doc, 1):
        lines = [[line] for line in page.get_text().splitlines() if line.strip()]
        if lines:
            yield f"Página {number}", lines


def _clean(cell):
    """Texto de celda sin saltos de línea internos; ``None`` queda como celda vacía."""
    return " ".join(cell.split()) if cell else ""


def _fit_columns(sheet):
    """Ajusta el ancho de cada columna a su texto más largo (con tope)."""
    for column in sheet.columns:
        longest = max((len(str(cell.value or "")) for cell in column), default=0)
        sheet.column_dimensions[column[0].column_letter].width = min(longest + 2, MAX_COLUMN_WIDTH)


def pdf_to_ppt(inputs, out, **_):
    """Crea una diapositiva por página con su imagen a tamaño completo.

    El texto de cada página se copia en las notas del orador para poder buscarlo
    o reutilizarlo; la diapositiva conserva el aspecto exacto del original.
    """
    from pptx import Presentation

    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    with open_pdf(inputs[0]) as doc:
        first = doc[0].rect
        presentation.slide_width = int(first.width * POINTS_TO_EMU)
        presentation.slide_height = int(first.height * POINTS_TO_EMU)
        for page in doc:
            slide = presentation.slides.add_slide(blank)
            image = page.get_pixmap(dpi=SLIDE_DPI, alpha=False).tobytes("jpg", jpg_quality=SLIDE_JPEG_QUALITY)
            slide.shapes.add_picture(BytesIO(image), 0, 0, presentation.slide_width, presentation.slide_height)
            slide.notes_slide.notes_text_frame.text = page.get_text().strip()
    result = out.with_suffix(".pptx")
    presentation.save(result)
    return result
