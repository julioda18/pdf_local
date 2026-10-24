import io
import sys
import zipfile
from pathlib import Path

import pikepdf
import pymupdf
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import app  # noqa: E402
from core.errors import ToolError  # noqa: E402
from core.params import parse_page_groups, parse_page_ranges  # noqa: E402
from tools import external, images, optimize, pages, security, stamps  # noqa: E402


@pytest.fixture
def make_pdf(tmp_path):
    def make(page_count=3, name="in.pdf"):
        doc = pymupdf.open()
        for n in range(1, page_count + 1):
            doc.new_page().insert_text((72, 72), f"Pagina {n}", fontsize=24)
        path = tmp_path / name
        doc.save(path)
        return path
    return make


def page_count(path):
    with pymupdf.open(path) as doc:
        return doc.page_count


def text_of(path, index=0):
    with pymupdf.open(path) as doc:
        return doc[index].get_text()


@pytest.mark.parametrize("spec, expected", [("1-3", [0, 1, 2]), ("2, 4", [1, 3]), ("4,1", [3, 0])])
def test_parse_page_ranges(spec, expected):
    assert parse_page_ranges(spec, 4) == expected


@pytest.mark.parametrize("spec", ["0", "1-9", "a", "3-1", ""])
def test_parse_page_ranges_rejects(spec):
    with pytest.raises(ToolError):
        parse_page_ranges(spec, 4)


def test_parse_page_groups_defaults_to_each_page():
    assert parse_page_groups("", 3) == [[0], [1], [2]]


def test_merge(make_pdf, tmp_path):
    assert page_count(pages.merge([make_pdf(2, "a.pdf"), make_pdf(3, "b.pdf")], tmp_path / "o.pdf")) == 5


def test_split(make_pdf, tmp_path):
    archive = zipfile.ZipFile(pages.split([make_pdf(5)], tmp_path / "o.pdf", ranges="1-2;3-5"))
    assert archive.namelist() == ["parte_1.pdf", "parte_2.pdf"]


def test_organize(make_pdf, tmp_path):
    out = pages.organize([make_pdf(4)], tmp_path / "o.pdf", order="4,1")
    assert page_count(out) == 2 and "Pagina 4" in text_of(out)


def test_rotate(make_pdf, tmp_path):
    with pymupdf.open(pages.rotate([make_pdf(2)], tmp_path / "o.pdf", angle="90", pages="2")) as doc:
        assert [p.rotation for p in doc] == [0, 90]


def test_rotate_rejects_bad_angle(make_pdf, tmp_path):
    with pytest.raises(ToolError):
        pages.rotate([make_pdf(1)], tmp_path / "o.pdf", angle="45")


def test_compress(make_pdf, tmp_path):
    assert page_count(optimize.compress([make_pdf(3)], tmp_path / "o.pdf", level="alta")) == 3


def test_protect_and_unlock(make_pdf, tmp_path):
    locked = security.protect([make_pdf(1)], tmp_path / "l.pdf", password="x1")
    with pytest.raises(pikepdf.PasswordError):
        pikepdf.open(locked)
    with pytest.raises(ToolError):
        security.unlock([locked], tmp_path / "u.pdf", password="mal")
    assert page_count(security.unlock([locked], tmp_path / "u.pdf", password="x1")) == 1


def test_locked_pdf_gives_clear_error(make_pdf, tmp_path):
    locked = security.protect([make_pdf(1)], tmp_path / "l.pdf", password="x1")
    with pytest.raises(ToolError, match="contraseña"):
        pages.merge([locked], tmp_path / "o.pdf")


def test_images_roundtrip(make_pdf, tmp_path):
    archive = zipfile.ZipFile(images.pdf_to_images([make_pdf(2)], tmp_path / "o.pdf", fmt="png", dpi="72"))
    extracted = []
    for name in archive.namelist():
        (tmp_path / name).write_bytes(archive.read(name))
        extracted.append(tmp_path / name)
    assert page_count(images.images_to_pdf(extracted, tmp_path / "i.pdf")) == 2


def test_watermark_and_page_numbers(make_pdf, tmp_path):
    marked = stamps.watermark([make_pdf(2)], tmp_path / "w.pdf", text="SECRETO")
    numbered = stamps.page_numbers([marked], tmp_path / "n.pdf")
    assert "SECRETO" in text_of(numbered) and "2 / 2" in text_of(numbered, 1)


def test_watermark_rejects_bad_opacity(make_pdf, tmp_path):
    with pytest.raises(ToolError):
        stamps.watermark([make_pdf(1)], tmp_path / "w.pdf", opacity="5")


@pytest.mark.skipif(not external.ocr_available(), reason="Tesseract/Ghostscript no instalados")
def test_ocr_makes_scanned_pdf_searchable(make_pdf, tmp_path):
    with pymupdf.open(make_pdf(1)) as source:
        pixmap, rect = source[0].get_pixmap(dpi=200), source[0].rect
    scanned = pymupdf.open()
    page = scanned.new_page(width=rect.width, height=rect.height)
    page.insert_image(page.rect, pixmap=pixmap)
    scanned.save(tmp_path / "scan.pdf")
    assert "Pagina" in text_of(external.ocr([tmp_path / "scan.pdf"], tmp_path / "o.pdf", lang="spa"))


@pytest.mark.skipif(not external.office_available(), reason="LibreOffice no instalado")
def test_office_to_pdf(tmp_path):
    source = tmp_path / "nota.txt"
    source.write_text("Hola desde LibreOffice", encoding="utf-8")
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    assert "Hola" in text_of(external.office_to_pdf([source], out_dir / "resultado.pdf"))


def test_api_rejects_fake_pdf():
    response = app.test_client().post("/api/merge", data={"files": (io.BytesIO(b"hola"), "x.pdf")})
    assert response.status_code == 400 and "PDF" in response.json["error"]


def test_api_merge(make_pdf):
    data = {"files": [(io.BytesIO(make_pdf(1, "a.pdf").read_bytes()), "a.pdf"),
                      (io.BytesIO(make_pdf(2, "b.pdf").read_bytes()), "b.pdf")]}
    response = app.test_client().post("/api/merge", data=data, content_type="multipart/form-data")
    assert response.status_code == 200 and pymupdf.open("pdf", response.data).page_count == 3


def test_preview_store_renders_on_demand(make_pdf):
    from tools.preview import PreviewStore
    store = PreviewStore()
    token, ratios = store.add(make_pdf(4))
    assert len(ratios) == 4 and ratios[0] > 1
    data, mimetype = store.thumbnail(token, 2)
    assert mimetype == "image/png" and data[1:4] == b"PNG"
    with pytest.raises(ToolError):
        store.thumbnail(token, 5)
    with pytest.raises(ToolError):
        store.thumbnail("caducado", 1)


def test_preview_store_evicts_oldest(make_pdf):
    from tools.preview import PreviewStore
    store = PreviewStore(max_documents=2)
    path = make_pdf(1)
    first, _ = store.add(path)
    store.add(path)
    store.add(path)
    with pytest.raises(ToolError):
        store.thumbnail(first, 1)


def test_preview_scanned_page_uses_jpeg(tmp_path):
    from tools.preview import render_page
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 200), False)
    pix.set_rect(pix.irect, (200, 100, 50))
    doc = pymupdf.open()
    doc.new_page().insert_image(pymupdf.Rect(0, 0, 595, 842), stream=pix.tobytes("png"))
    assert render_page(doc[0])[1] == "image/jpeg"


def test_preview_endpoints(make_pdf):
    client = app.test_client()
    with make_pdf(3).open("rb") as f:
        body = client.post("/api/preview", data={"files": (f, "a.pdf")}).get_json()
    assert len(body["ratios"]) == 3
    page = client.get(f"/api/preview/{body['id']}/3")
    assert page.status_code == 200 and page.mimetype == "image/png" and "max-age" in page.headers["Cache-Control"]
    assert client.get(f"/api/preview/{body['id']}/9").status_code == 400
    bad = client.post("/api/preview", data={"files": (io.BytesIO(b"no pdf"), "a.pdf")})
    assert bad.status_code == 400 and "error" in bad.get_json()


@pytest.mark.parametrize("ranges, parts", [("1-2;3-5", 2), ("1-1;2-2;3-5", 3)])
def test_split_ui_ranges(make_pdf, tmp_path, ranges, parts):
    with zipfile.ZipFile(pages.split([make_pdf(5)], tmp_path / "o.pdf", ranges=ranges)) as archive:
        assert len(archive.namelist()) == parts


def test_pdf_to_word(make_pdf, tmp_path):
    from docx import Document
    from tools import export
    result = export.pdf_to_word([make_pdf(2)], tmp_path / "out.pdf")
    assert result.suffix == ".docx"
    assert "Pagina 2" in "\n".join(p.text for p in Document(result).paragraphs)


def test_pdf_to_word_endpoint(make_pdf):
    with make_pdf(1).open("rb") as f:
        response = app.test_client().post("/api/pdf_to_word", data={"files": (f, "a.pdf")})
    assert response.status_code == 200 and response.data[:2] == b"PK"
    assert "a.docx" in response.headers["Content-Disposition"]


@pytest.fixture
def table_pdf(tmp_path):
    doc = pymupdf.open()
    page = doc.new_page()
    rows = [("Producto", "Precio"), ("Manzana", "10"), ("Pera", "12")]
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            rect = pymupdf.Rect(72 + c * 150, 72 + r * 30, 222 + c * 150, 102 + r * 30)
            page.draw_rect(rect, color=(0, 0, 0))
            page.insert_text((rect.x0 + 5, rect.y1 - 10), text)
    path = tmp_path / "tabla.pdf"
    doc.save(path)
    return path


def test_pdf_to_excel_tables(table_pdf, tmp_path):
    from openpyxl import load_workbook
    from tools import export
    sheet = load_workbook(export.pdf_to_excel([table_pdf], tmp_path / "o.pdf")).worksheets[0]
    assert sheet.title.startswith("Pág 1") and [c.value for c in sheet[2]] == ["Manzana", "10"]


def test_pdf_to_excel_text_fallback(make_pdf, tmp_path):
    from openpyxl import load_workbook
    from tools import export
    book = load_workbook(export.pdf_to_excel([make_pdf(2)], tmp_path / "o.pdf"))
    assert book.sheetnames == ["Página 1", "Página 2"] and book.worksheets[1]["A1"].value == "Pagina 2"


def test_pdf_to_excel_without_text(tmp_path):
    from tools import export
    path = tmp_path / "vacio.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(path)
    with pytest.raises(ToolError):
        export.pdf_to_excel([path], tmp_path / "o.pdf")


def test_pdf_to_ppt(make_pdf, tmp_path):
    from pptx import Presentation
    from tools import export
    deck = Presentation(export.pdf_to_ppt([make_pdf(3)], tmp_path / "o.pdf"))
    assert len(deck.slides) == 3
    assert "Pagina 3" in deck.slides[2].notes_slide.notes_text_frame.text


@pytest.mark.parametrize("tool, suffix", [("pdf_to_excel", ".xlsx"), ("pdf_to_ppt", ".pptx")])
def test_office_export_endpoints(make_pdf, tool, suffix):
    with make_pdf(1).open("rb") as f:
        response = app.test_client().post(f"/api/{tool}", data={"files": (f, "a.pdf")})
    assert response.status_code == 200 and response.data[:2] == b"PK"
    assert f"a{suffix}" in response.headers["Content-Disposition"]
